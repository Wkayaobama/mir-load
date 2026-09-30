#!/usr/bin/env python3
"""Bootstrap the spec from an existing execution-surfaces workbook: .xlsx → YAML.

    python3 dump_surfaces_xlsx.py --xlsx docs/EXECUTION_SURFACES.xlsx --out docs/execution_surfaces.yaml

Use it once, when a workbook was built by hand (or by an earlier session) and the YAML source does not exist
yet. From then on edit the YAML and rebuild with build_surfaces_xlsx.py; the round trip is checked with
`build_surfaces_xlsx.py --compare <original.xlsx>`. Expects the five sheets and column layouts the producer
writes (Overview / Layers / Pipeline steps / Remediation / Equality checks).
"""
from __future__ import annotations

import argparse
import sys
from pathlib import Path

import openpyxl
import yaml


class _Str(str):
    pass


def _lit(dumper, data):  # long text as block scalars → readable diffs
    style = "|" if ("\n" in data or len(data) > 80) else None
    return dumper.represent_scalar("tag:yaml.org,2002:str", data, style=style)


yaml.add_representer(str, _lit)


def _rows(ws, first_row: int, ncols: int, stop_at_blank=True):
    out = []
    for r in range(first_row, ws.max_row + 1):
        vals = [ws.cell(r, c).value for c in range(1, ncols + 1)]
        if stop_at_blank and vals[0] in (None, ""):
            break
        out.append(["" if v is None else v for v in vals])
    return out


def dump(path: Path) -> dict:
    wb = openpyxl.load_workbook(path)
    ov, ly, ps, rm, eq = (wb[n] for n in ("Overview", "Layers", "Pipeline steps", "Remediation", "Equality checks"))
    motions = []
    for r in range(5, ov.max_row + 1):
        a = ov.cell(r, 1).value
        if a in (None, "") or str(a).startswith("Parity summary"):
            break
        motions.append(dict(zip(("motion", "acquire", "ubuntu", "driver", "state", "assessment"),
                                ("" if ov.cell(r, c).value is None else ov.cell(r, c).value for c in range(1, 7)))))
    legend = next((ov.cell(r, 1).value for r in range(ov.max_row, 0, -1) if str(ov.cell(r, 1).value or "").startswith("Legend")), "")
    spec = {
        "title": ov["A1"].value or "",
        "verdict": ov["A2"].value or "",
        "legend": legend,
        "surfaces": {"A": ly["C1"].value, "B": ly["D1"].value, "C": ly["E1"].value},
        "motions": motions,
        "layers": [dict(zip(("id", "layer", "A", "B", "C", "parity", "diverges", "remediation", "verified"), v)) for v in _rows(ly, 2, 9)],
        "steps": [dict(zip(("n", "step", "gate", "A", "B", "C", "parity", "artefact", "notes"), v)) for v in _rows(ps, 2, 9)],
        "remediation": [dict(zip(("id", "symptom", "layer", "surfaces", "cause", "fix", "verified"), v)) for v in _rows(rm, 2, 7)],
        "equality_checks": [dict(zip(("stage", "compare", "A", "BC", "expected"), v)) for v in _rows(eq, 2, 5)],
    }
    return spec


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--xlsx", required=True)
    ap.add_argument("--out", required=True, help="YAML spec to write")
    a = ap.parse_args(argv)
    spec = dump(Path(a.xlsx))
    Path(a.out).write_text(yaml.dump(spec, sort_keys=False, allow_unicode=True, width=1000), encoding="utf-8")
    print(f"wrote {a.out}: {len(spec['layers'])} layers, {len(spec['steps'])} steps, {len(spec['remediation'])} remediations, "
          f"{len(spec['equality_checks'])} equality checks, {len(spec['motions'])} motions")
    return 0


if __name__ == "__main__":
    sys.exit(main())
