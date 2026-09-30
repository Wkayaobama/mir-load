#!/usr/bin/env python3
"""Execution-surfaces workbook producer: spec (YAML) → .xlsx, with the conventions of the mr-load workbook.

    python3 build_surfaces_xlsx.py --spec docs/execution_surfaces.yaml --out docs/EXECUTION_SURFACES.xlsx --verify
    python3 build_surfaces_xlsx.py --spec … --out … --compare docs/EXECUTION_SURFACES.xlsx   # value diff vs an existing workbook

Sheets: Overview (motions, live parity formulas, legend) · Layers · Pipeline steps · Remediation · Equality checks.
Conventions: Arial 10, navy header row (white bold), wrapped/top-aligned body, thin borders, parity colours
(green Parallel, yellow Conditional, orange Divergent), frozen header rows, fullCalcOnLoad so the COUNTIF /
SUM / COUNTA cells compute when the file is opened in Excel or LibreOffice.

--verify reloads the file and checks what LibreOffice would have shown: every parity value is one of the three
words, the parity counts add up to the row counts the COUNTA formulas will return, fonts are Arial, formulas
present. Use it whenever LibreOffice recalculation is unavailable (the workbook then recalculates on open).
Exit code 1 on a failed verification or a non-empty comparison.
"""
from __future__ import annotations

import argparse
import math
import sys
from copy import copy
from pathlib import Path

import openpyxl
import yaml
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

HEADER_FILL = "001F3864"
PARITY_FILL = {"Parallel": "00E2F0D9", "Conditional": "00FFF2CC", "Divergent": "00F8CBAD"}
PARITY_MEANING = {
    "Parallel": "same mechanism, same result on every surface",
    "Conditional": "same result when the stated condition holds (same commit, same identity, same clone for the ledger)",
    "Divergent": "different mechanism; the remediation makes it equivalent",
}
THIN = Side(style="thin", color="00BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
FONT = Font(name="Arial", size=10)
FONT_B = Font(name="Arial", size=10, bold=True)
FONT_H = Font(name="Arial", size=10, bold=True, color="00FFFFFF")
WRAP = Alignment(wrap_text=True, vertical="top")

WIDTHS = {
    "Overview": [40, 40, 34, 44, 34, 48],
    "Layers": [5, 24, 46, 40, 34, 12, 44, 52, 34],
    "Pipeline steps": [5, 22, 20, 40, 40, 40, 12, 50, 50],
    "Remediation": [6, 48, 8, 16, 46, 60, 24],
    "Equality checks": [18, 26, 70, 60, 44],
}


def _row_height(text: str) -> int:
    return 15 + 15 * max(1, math.ceil(len(text or "") / 150))


def _widths(ws, widths):
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def _header(ws, row, values):
    for i, v in enumerate(values, start=1):
        c = ws.cell(row=row, column=i, value=v)
        c.font = FONT_H; c.fill = PatternFill("solid", fgColor=HEADER_FILL); c.alignment = WRAP; c.border = BORDER


def _body(ws, row, values, parity_col=None):
    for i, v in enumerate(values, start=1):
        c = ws.cell(row=row, column=i, value=v)
        c.font = FONT; c.alignment = WRAP; c.border = BORDER
        if parity_col and i == parity_col:
            if v not in PARITY_FILL:
                raise SystemExit(f"{ws.title} row {row}: parity must be one of {list(PARITY_FILL)}, got {v!r}")
            c.font = FONT_B; c.fill = PatternFill("solid", fgColor=PARITY_FILL[v])


def build(spec: dict) -> openpyxl.Workbook:
    s = {k: v for k, v in spec.items()}
    surf = s["surfaces"]
    wb = openpyxl.Workbook()

    # ---- Overview
    ov = wb.active; ov.title = "Overview"; _widths(ov, WIDTHS["Overview"])
    ov["A1"] = s["title"]; ov["A1"].font = Font(name="Arial", size=14, bold=True)
    ov["A2"] = s["verdict"]; ov["A2"].font = FONT; ov["A2"].alignment = WRAP
    ov.merge_cells("A2:F2"); ov.row_dimensions[2].height = _row_height(s["verdict"])
    _header(ov, 4, ["Motion", "Acquire", "Ubuntu", "Driver", "State lives in", "DevOps assessment"])
    r = 5
    for m in s["motions"]:
        _body(ov, r, [m["motion"], m["acquire"], m["ubuntu"], m["driver"], m["state"], m["assessment"]]); r += 1
    r += 1
    ov.cell(row=r, column=1, value="Parity summary (live formulas over the other sheets)").font = FONT_B; r += 1
    _header(ov, r, ["Status", "Layers", "Pipeline steps", "Meaning"]); r += 1
    first = r
    for status in ("Parallel", "Conditional", "Divergent"):
        _body(ov, r, [status, f'=COUNTIF(Layers!$F:$F,"{status}")', f"=COUNTIF('Pipeline steps'!$G:$G,\"{status}\")",
                      s.get("parity_meanings", PARITY_MEANING)[status]], parity_col=1)
        r += 1
    _body(ov, r, ["Total", f"=SUM(B{first}:B{r-1})", f"=SUM(C{first}:C{r-1})", ""]); ov.cell(row=r, column=1).font = FONT_B; r += 1
    _body(ov, r, ["Rows on the sheet (must equal Total)", "=COUNTA(Layers!$A:$A)-1", "=COUNTA('Pipeline steps'!$A:$A)-1", ""]); r += 2
    ov.cell(row=r, column=1, value=s["legend"]).font = Font(name="Arial", size=9, color="00555555")
    ov.cell(row=r, column=1).alignment = WRAP; ov.merge_cells(f"A{r}:F{r}"); ov.row_dimensions[r].height = _row_height(s["legend"])
    ov.freeze_panes = "A5"

    # ---- Layers
    ly = wb.create_sheet("Layers"); _widths(ly, WIDTHS["Layers"])
    _header(ly, 1, ["#", "Layer", surf["A"], surf["B"], surf["C"], "Parity", "Where it diverges", "Remediation", "Verified how"])
    for i, L in enumerate(s["layers"], start=2):
        _body(ly, i, [L["id"], L["layer"], L["A"], L["B"], L["C"], L["parity"], L["diverges"], L["remediation"], L["verified"]], parity_col=6)
    ly.freeze_panes = "A2"

    # ---- Pipeline steps
    ps = wb.create_sheet("Pipeline steps"); _widths(ps, WIDTHS["Pipeline steps"])
    _header(ps, 1, ["#", "Step", "Gate", f"{surf['A']} — cell", f"{surf['B']} — command", f"{surf['C']} — command",
                    "Parity", "Identical artefact to compare", "Notes"])
    for i, st in enumerate(s["steps"], start=2):
        _body(ps, i, [st["n"], st["step"], st.get("gate", ""), st["A"], st["B"], st["C"], st["parity"],
                      st.get("artefact", ""), st.get("notes", "")], parity_col=7)
    ps.freeze_panes = "A2"

    # ---- Remediation
    rm = wb.create_sheet("Remediation"); _widths(rm, WIDTHS["Remediation"])
    _header(rm, 1, ["ID", "Symptom", "Layer", "Surface(s)", "Cause", "Fix", "Verified / source"])
    for i, R in enumerate(s["remediation"], start=2):
        _body(rm, i, [R["id"], R["symptom"], R["layer"], R["surfaces"], R["cause"], R["fix"], R["verified"]])
    rm.freeze_panes = "A2"

    # ---- Equality checks
    eq = wb.create_sheet("Equality checks"); _widths(eq, WIDTHS["Equality checks"])
    _header(eq, 1, ["Stage", "What to compare", f"{surf['A']} — cell", f"{surf['B']} / {surf['C']} — command", "Expected"])
    for i, E in enumerate(s["equality_checks"], start=2):
        _body(eq, i, [E["stage"], E["compare"], E["A"], E["BC"], E["expected"]])
    eq.freeze_panes = "A2"

    wb.calculation.fullCalcOnLoad = True
    return wb


def verify(path: Path) -> list[str]:
    """What a recalculation would show, checked without LibreOffice. Returns problems (empty = OK)."""
    wb = openpyxl.load_workbook(path)
    problems: list[str] = []
    for name, col in (("Layers", 6), ("Pipeline steps", 7)):
        ws = wb[name]
        counts = {k: 0 for k in PARITY_FILL}
        for r in range(2, ws.max_row + 1):
            v = ws.cell(r, col).value
            if v not in counts:
                problems.append(f"{name} row {r}: parity {v!r}")
            else:
                counts[v] += 1
        if sum(counts.values()) != ws.max_row - 1:
            problems.append(f"{name}: parity total {sum(counts.values())} != rows {ws.max_row - 1}")
        print(f"{name}: {counts}  rows={ws.max_row - 1}  (what the Overview formulas will show)")
    for ws in wb.worksheets:
        fonts = {ws.cell(r, c).font.name for r in range(1, ws.max_row + 1) for c in range(1, ws.max_column + 1)
                 if ws.cell(r, c).value is not None}
        if fonts - {"Arial"}:
            problems.append(f"{ws.title}: fonts {fonts}")
    formulas = [c.value for row in wb["Overview"].iter_rows() for c in row if isinstance(c.value, str) and c.value.startswith("=")]
    if len(formulas) != 10:
        problems.append(f"Overview: expected 10 formulas, found {len(formulas)}")
    if not wb.calculation.fullCalcOnLoad:
        problems.append("fullCalcOnLoad is off: formulas would show empty until a manual recalculation")
    return problems


def compare(a: Path, b: Path) -> list[str]:
    """Cell-by-cell value differences between two workbooks (styles ignored)."""
    wa, wb_ = openpyxl.load_workbook(a), openpyxl.load_workbook(b)
    diffs: list[str] = []
    if wa.sheetnames != wb_.sheetnames:
        diffs.append(f"sheets {wa.sheetnames} != {wb_.sheetnames}")
    for name in wa.sheetnames:
        if name not in wb_.sheetnames:
            continue
        sa, sb = wa[name], wb_[name]
        rows, cols = max(sa.max_row, sb.max_row), max(sa.max_column, sb.max_column)
        for r in range(1, rows + 1):
            for c in range(1, cols + 1):
                va, vb = sa.cell(r, c).value, sb.cell(r, c).value
                if (va or None) != (vb or None):
                    diffs.append(f"{name}!{get_column_letter(c)}{r}: {str(va)[:60]!r} != {str(vb)[:60]!r}")
    return diffs


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--spec", required=True, help="YAML spec (see references/surfaces-spec.md)")
    ap.add_argument("--out", required=True, help="workbook to write")
    ap.add_argument("--verify", action="store_true", help="reload and check parity counts / fonts / formulas")
    ap.add_argument("--compare", help="existing workbook to diff values against after building")
    a = ap.parse_args(argv)
    spec = yaml.safe_load(Path(a.spec).read_text(encoding="utf-8"))
    out = Path(a.out); out.parent.mkdir(parents=True, exist_ok=True)
    build(spec).save(out)
    print(f"wrote {out}")
    rc = 0
    if a.verify:
        problems = verify(out)
        print("verify:", "OK" if not problems else "\n  " + "\n  ".join(problems))
        rc |= bool(problems)
    if a.compare:
        diffs = compare(out, Path(a.compare))
        print(f"compare vs {a.compare}:", "identical values" if not diffs else f"{len(diffs)} differences\n  " + "\n  ".join(diffs[:40]))
        rc |= bool(diffs)
    return rc


if __name__ == "__main__":
    sys.exit(main())
